import io

import pandas as pd
from openpyxl import load_workbook

from conftest import ADMIN, APPROVER, SUBMITTER


def _ids(client):
    periods = {p["code"]: p["id"] for p in client.get("/api/periods", headers=ADMIN).json()}
    bus = {b["code"]: b["id"] for b in client.get("/api/business-units", headers=ADMIN).json()}
    kpis = {k["code"]: k["id"] for k in client.get("/api/kpis", headers=ADMIN).json()}
    return periods, bus, kpis


def test_auth_and_roles(client):
    assert client.get("/api/kpis").status_code == 401
    r = client.post("/api/business-units", json={"code": "X", "name": "X"}, headers=SUBMITTER)
    assert r.status_code == 403


def test_kpi_formula_requires_approval(client):
    r = client.post("/api/kpis", headers=ADMIN, json={
        "code": "REV_PER_TICKET", "name": "Revenue per ticket", "kpi_type": "formula",
        "formula": "REVENUE / TICKETS_CLOSED", "formula_approved": True})
    assert r.status_code == 201 and r.json()["formula_approved"] is False
    bad = client.post("/api/kpis", headers=ADMIN, json={"code": "BAD", "name": "b", "kpi_type": "formula",
                                                         "formula": "__import__('os').system('x')"})
    assert bad.status_code == 422
    selfref = client.post("/api/kpis", headers=ADMIN, json={"code": "SELF", "name": "s", "kpi_type": "formula",
                                                            "formula": "SELF + 1"})
    assert selfref.status_code == 422
    ok = client.post(f"/api/kpis/{r.json()['id']}/approve-formula", headers=APPROVER)
    assert ok.json()["formula_approved"] is True


def test_manual_entry_submit_and_segregation(client):
    periods, bus, kpis = _ids(client)
    body = {"period_id": periods["2026-09"], "business_unit_id": bus["RETAIL"], "action": "submit",
            "entries": [{"kpi_id": kpis["REVENUE"], "actual": 130, "target": 125, "commentary": "Festive uplift"},
                        {"kpi_id": kpis["OPEX"], "actual": 66, "target": 68}]}
    r = client.put("/api/values", json=body, headers=SUBMITTER)
    assert r.status_code == 200, r.text
    vid = r.json()[0]["id"]
    # locked after submit
    body["entries"][0]["actual"] = 999
    assert client.put("/api/values", json=body, headers=SUBMITTER).status_code == 422
    # submitter-as-approver blocked
    self_review = client.post(f"/api/values/{vid}/review", json={"approve": True},
                              headers={"X-User": "sam", "X-Role": "approver"})
    assert self_review.status_code == 403
    # evidence upload
    ev = client.post(f"/api/values/{vid}/evidence", headers=SUBMITTER,
                     files={"file": ("proof.txt", b"ledger extract", "text/plain")})
    assert ev.status_code == 201
    grid = client.get(f"/api/values/grid?period_id={periods['2026-09']}&business_unit_id={bus['RETAIL']}",
                      headers=SUBMITTER).json()
    assert any(r["evidence_count"] == 1 for r in grid["rows"])


def test_file_submission_validation(client):
    periods, bus, kpis = _ids(client)
    tpl = client.get("/api/files/template", headers=SUBMITTER)
    assert tpl.status_code == 200 and "Submission" in load_workbook(io.BytesIO(tpl.content)).sheetnames

    df = pd.DataFrame([
        ["REVENUE", "CORP", "2026-09", "210", "200", "Large deal closed"],
        ["OPEX", "CORP", "2026-09", "1,050", "1100", ""],             # thousands separator accepted
        ["TICKETS_CLOSED", "CORP", "2026-09", "1500", "1520", ""],
        ["TICKETS_OPENED", "CORP", "2026-09", "1600", "1600", ""],
        ["CSAT", "CORP", "2026-09", "86", "85", ""],
        ["NOPE", "CORP", "2026-09", "1", "", ""],                     # unknown KPI
        ["REVENUE", "MARS", "2026-09", "1", "", ""],                  # unknown BU
        ["COST_INCOME", "CORP", "2026-09", "50", "", ""],             # formula KPI
        ["REVENUE", "CORP", "2026-09", "220", "", ""],                # duplicate
        ["REVENUE", "OPS", "2026-08", "50", "", ""],                  # closed period
        ["REVENUE", "OPS", "2026-09", "abc", "", ""],                 # non-numeric
    ], columns=["KPI Code", "Business Unit", "Period", "Actual", "Target", "Commentary"])
    buf = io.BytesIO()
    df.to_excel(buf, index=False, sheet_name="Submission")
    r = client.post("/api/files/submissions", headers=SUBMITTER,
                    files={"file": ("corp_sep.xlsx", buf.getvalue(), "application/octet-stream")})
    job = r.json()
    assert job["status"] == "completed", job
    assert (job["rows_total"], job["rows_accepted"], job["rows_rejected"]) == (11, 5, 6)
    rej = client.get(f"/api/jobs/{job['id']}/rejected", headers=SUBMITTER).json()
    reasons = " ".join(" ".join(x["reasons"]) for x in rej)
    for frag in ["Unknown or inactive KPI", "Unknown or inactive business unit", "calculated by formula",
                 "Duplicate", "closed", "not a valid number"]:
        assert frag in reasons
    assert client.get(f"/api/jobs/{job['id']}/rejected.xlsx", headers=SUBMITTER).status_code == 200
    assert client.get(f"/api/jobs/{job['id']}/logs", headers=SUBMITTER).json()

    bad = client.post("/api/files/submissions", headers=SUBMITTER,
                      files={"file": ("x.csv", b"foo,bar\n1,2\n", "text/csv")})
    assert bad.json()["status"] == "failed" and "missing columns" in bad.json()["error"]


def test_approve_consolidate_export(client):
    periods, bus, kpis = _ids(client)
    pending = client.get("/api/values/pending", headers=APPROVER).json()
    assert len(pending) == 7
    r = client.post("/api/values/review-bulk", headers=APPROVER,
                    json={"value_ids": [p["value_id"] for p in pending], "approve": True})
    assert r.json()["reviewed"] == 7

    job = client.post("/api/consolidations", headers=APPROVER, json={"period_id": periods["2026-09"]}).json()
    assert job["status"] == "completed", job
    data = client.get(f"/api/consolidations/latest?period_id={periods['2026-09']}", headers=SUBMITTER).json()
    rows = {(r["business_unit"], r["kpi_code"]): r for r in data["rows"]}
    # deterministic checks
    assert rows[("GROUP", "REVENUE")]["actual"] == 340            # 130 + 210
    assert rows[("CORP", "CLOSURE_RATE")]["actual"] == 93.75      # 1500/1600*100
    assert rows[("RETAIL", "COST_INCOME")]["actual"] == round(66 / 130 * 100, 4)
    assert rows[("CORP", "REVENUE")]["rag"] == "GREEN"
    assert rows[("RETAIL", "OPEX")]["rag"] == "GREEN"             # lower is better
    assert rows[("CORP", "REVENUE")]["prior_actual"] == round(120 * 1.6 * 0.94, 2)
    assert rows[("CORP", "REVENUE")]["change_vs_prior_pct"] == round((210 - 180.48) / 180.48 * 100, 2)
    xl = client.get(f"/api/documents/{job['result']['excel_document_id']}/download", headers=SUBMITTER)
    wb = load_workbook(io.BytesIO(xl.content))
    assert {"Group Summary", "Exceptions", "Consolidated Dataset"} <= set(wb.sheetnames)
    pp = client.get(f"/api/documents/{job['result']['pptx_document_id']}/download", headers=SUBMITTER)
    assert pp.status_code == 200 and len(pp.content) > 10000
    trend = client.get("/api/dashboard/trend?kpi_code=REVENUE&business_unit=GROUP", headers=SUBMITTER).json()
    assert [p["period"] for p in trend["points"]] == ["2026-07", "2026-08", "2026-09"]
    status = client.get(f"/api/dashboard/status?period_id={periods['2026-09']}", headers=SUBMITTER).json()
    assert {s["business_unit"]: s["approved"] for s in status}["CORP"] == 5


def test_commentary_workflow_and_ask(client):
    periods, _, _ = _ids(client)
    pid = periods["2026-09"]
    c = client.post("/api/commentary", headers=SUBMITTER, json={"period_id": pid, "business_unit": "GROUP"}).json()
    assert c["status"] == "draft" and "Exceptions" in c["draft_text"]
    # cannot publish before review
    assert client.post(f"/api/commentary/{c['id']}/publish", headers=APPROVER).status_code == 409
    edited = client.put(f"/api/commentary/{c['id']}", headers=SUBMITTER,
                        json={"draft_text": c["draft_text"] + "\nRevenue reached 98765.43 this month."}).json()
    assert any("98765.43" in w for w in edited["validation_warnings"])
    client.post(f"/api/commentary/{c['id']}/submit", headers=SUBMITTER)
    # requester cannot review own commentary
    assert client.post(f"/api/commentary/{c['id']}/review", json={"approve": True},
                       headers={"X-User": "sam", "X-Role": "approver"}).status_code == 403
    assert client.post(f"/api/commentary/{c['id']}/review", json={"approve": False},
                       headers=APPROVER).status_code == 422  # notes required
    ok = client.post(f"/api/commentary/{c['id']}/review", json={"approve": True, "final_text": "Final text"},
                     headers=APPROVER).json()
    assert ok["status"] == "approved"
    pub = client.post(f"/api/commentary/{c['id']}/publish", headers=APPROVER).json()
    assert pub["status"] == "published" and pub["final_text"] == "Final text"

    doc = client.post("/api/documents", headers=SUBMITTER, files={"file": (
        "kpi_policy.txt", b"Cost-to-income ratio measures operating expenses divided by revenue. "
                          b"Corporate banking targets are set annually by the CFO office.", "text/plain")})
    assert doc.status_code == 201
    docs = client.get("/api/documents?doc_type=reference", headers=SUBMITTER).json()
    assert docs[0]["ingested"] is True
    ans = client.post("/api/ask", headers=SUBMITTER,
                      json={"question": "What is the cost-to-income ratio for CORP?", "period_id": pid}).json()
    assert "Cost-to-Income Ratio" in ans["answer"] and ans["sources"][0]["document"] == "kpi_policy.txt"
