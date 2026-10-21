"""Tracker import in the team's real layout and the monthly MIS summary."""
import io
from datetime import date

from openpyxl import Workbook, load_workbook

from app.services.demo_tracker import build_sample_tracker, demo_records
from app.services.mis_summary import movement
from app.services.tracker_import import match_columns, normalise_type, read_table, split_names
from conftest import SUBMITTER


def _expected(records, y, m):
    rows = [r for r in records if r["actual_date"] and (r["actual_date"].year, r["actual_date"].month) == (y, m)]
    by_type = {t: sum(1 for r in rows if r["project_type"] == t) for t in ("Stat", "IA", "Tax")}
    kpis = sum(r["kpis_delivered"] or 0 for r in rows)
    return rows, by_type, kpis


def test_unit_helpers():
    assert normalise_type("Statutory") == "Stat" and normalise_type("Internal Audit") == "IA"
    assert normalise_type(" ia ") == "IA" and normalise_type("TAX") == "Tax" and normalise_type(None) == "Other"
    assert split_names("akhilesh, Isha & rahul") == ["Akhilesh", "Isha", "Rahul"]
    assert movement(58, 48)["text"] == "Up by 21%" and movement(899, 631)["text"] == "Up by 42%"
    assert movement(25, 44)["direction"] == "down" and movement(3, 0)["pct"] is None


def test_reads_real_layout_headers():
    table = read_table(build_sample_tracker(), "tracker.xlsx")
    assert table["header_row"] == 13  # legend rows above are skipped
    texts = [h["text"] for h in table["headers"]]
    assert "Step 4 - KPI Execution Number of KPIs Delivered" in texts  # merged parent + sub header
    mapping = match_columns(table["headers"])
    pick = {f: texts[i] for f, i in mapping.items() if i is not None}
    assert pick["kpis_delivered"].endswith("Number of KPIs Delivered")
    assert pick["kpis_planned"].endswith("Number of KPIs in scope")
    assert pick["actual_date"] == "Actual Date of Delivery" and pick["expected_date"] == "Expected Date of Delivery"
    assert pick["business_area"] == "Business area" and pick["delivery_lead"] == "Delivery Lead"


def test_import_and_summary_match_source(client):
    records = demo_records()
    r = client.post("/api/tracker/import", headers=SUBMITTER, data={"mode": "replace"},
                    files={"file": ("Project tracker.xlsx", build_sample_tracker(records), "application/octet-stream")})
    body = r.json()
    assert body["status"] == "completed", body
    assert body["engagements"] == len(records) and body["missing"] == []

    s = client.get("/api/mis/summary?month=2026-10&compare=py", headers=SUBMITTER).json()
    rows, by_type, kpis = _expected(records, 2026, 10)
    _, prev_type, prev_kpis = _expected(records, 2025, 10)
    eng, kp = s["headline"]
    assert eng["current"]["Total"] == len(rows) and {t: eng["current"][t] for t in by_type} == by_type
    assert eng["compare"]["Total"] == sum(prev_type.values()) and kp["current"]["Total"] == kpis
    assert kp["compare"]["Total"] == prev_kpis
    assert s["compare_label"] == "Oct-25 (Same month – PY)"
    assert sum(x["engagements"] for x in s["regions"]) == len(rows)
    assert [x["region"] for x in s["regions"]][:4] == ["North", "South", "East", "West"]
    assert s["resource_totals"]["lead"] == sum(1 for r in rows if r["delivery_lead"])
    assert s["resource_totals"]["secondary"] == sum(1 for r in rows if r["secondary_resource"])
    assert sum(p["engagements"] for p in s["partners"]) == len(rows)
    assert len(s["trend"]) == 14 and s["trend"][-1]["label"] == "Oct-26" and s["trend"][0]["label"] == "Sep-25"
    assert s["trend"][-1]["engagements"] == len(rows)

    pm = client.get("/api/mis/summary?month=2026-10&compare=pm", headers=SUBMITTER).json()
    sep, _, _ = _expected(records, 2026, 9)
    assert pm["headline"][0]["compare"]["Total"] == len(sep) and pm["compare_label"].startswith("Sep-26")

    xl = client.get("/api/mis/summary.xlsx?month=2026-10&compare=py", headers=SUBMITTER)
    wb = load_workbook(io.BytesIO(xl.content))
    assert {"MIS Summary", "Chart data", "Delivered engagements"} <= set(wb.sheetnames)
    assert len(wb["MIS Summary"]._charts) == 2
    assert wb["Delivered engagements"].max_row == len(rows) + 1
    lst = client.get("/api/mis/engagements?month=2026-10", headers=SUBMITTER).json()
    assert len(lst) == len(rows)


def test_messy_tracker_and_remap(client):
    """Odd header names, text dates, blanks; user fixes the column match and it is remembered."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Monthly tracker - do not edit legend"])
    ws.append([])
    ws.append(["S.No", "Project Name", "Engagement Kind", "KPIs done", "Completed on", "EP", "Zone"])
    ws.append([1, "Alpha", "Statutory", 10, "05/10/2026", "P One", "north"])
    ws.append([2, "Beta", "Internal", "12", "2026-10-07", "P Two", "west"])
    ws.append([3, "Gamma", "", "", "", "P One", ""])          # not delivered
    ws.append([4, "Delta", "Stat", "abc", "31-10-2026", None, "North"])
    ws.append([None, "Total", None, 22, None, None, None])     # totals row ignored
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/tracker/import", headers=SUBMITTER, data={"mode": "replace"},
                    files={"file": ("odd.xlsx", buf.getvalue(), "application/octet-stream")}).json()
    assert r["engagements"] == 4
    assert set(r["missing"]) >= {"project_type", "kpis_delivered", "actual_date"}
    idx = {h["text"]: h["index"] for h in r["headers"]}
    fixed = client.post(f"/api/tracker/import/{r['id']}/remap", headers=SUBMITTER, json={"mapping": {
        "project_type": idx["Engagement Kind"], "kpis_delivered": idx["KPIs done"], "actual_date": idx["Completed on"],
        "partner": idx["EP"], "region": idx["Zone"]}}).json()
    assert fixed["status"] == "completed" and fixed["missing"] == [] and fixed["delivered"] == 3
    assert any("not a valid number" in k for k in fixed["warnings"])
    s = client.get("/api/mis/summary?month=2026-10&compare=pm", headers=SUBMITTER).json()
    assert s["headline"][0]["current"] == {"Stat": 2, "IA": 1, "Tax": 0, "Total": 3}
    assert s["headline"][1]["current"]["Total"] == 22  # 10 + 12 + blank
    assert {x["region"]: x["engagements"] for x in s["regions"]}["North"] == 2

    # same layout again: remembered matches apply without remapping
    again = client.post("/api/tracker/import", headers=SUBMITTER, data={"mode": "replace"},
                        files={"file": ("odd.xlsx", buf.getvalue(), "application/octet-stream")}).json()
    assert again["missing"] == []

    # append last year's file without losing this year's rows
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.append(["S.No", "Project Name", "Engagement Kind", "KPIs done", "Completed on", "EP", "Zone"])
    ws2.append([1, "Old one", "Stat", 7, date(2025, 10, 3), "P One", "South"])
    buf2 = io.BytesIO()
    wb2.save(buf2)
    client.post("/api/tracker/import", headers=SUBMITTER, data={"mode": "append"},
                files={"file": ("ly.xlsx", buf2.getvalue(), "application/octet-stream")})
    py = client.get("/api/mis/summary?month=2026-10&compare=py", headers=SUBMITTER).json()
    assert py["headline"][0]["current"]["Total"] == 3 and py["headline"][0]["compare"]["Total"] == 1
    assert py["headline"][0]["movement"]["text"] == "Up by 200%"
