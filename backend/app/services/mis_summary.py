"""Monthly MIS summary from imported tracker engagements.

An engagement counts as delivered in the month of its actual delivery date. KPIs executed = the
number of KPIs delivered on those engagements. The comparison month is either the previous month
or the same month of the previous year.
"""
import io
import math
from calendar import monthrange
from collections import Counter
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..models import Engagement, ProcessingJob
from .tracker_import import FIELDS, PROJECT_TYPES, split_names

STANDARD_REGIONS = ["North", "South", "East", "West"]
NOT_SPECIFIED = "Not specified"
MONTH_ABBR = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def parse_month(s: str) -> tuple[int, int]:
    y, m = s.split("-")
    y, m = int(y), int(m)
    if not 1 <= m <= 12:
        raise ValueError("month must be 1-12")
    return y, m


def shift(y: int, m: int, months: int) -> tuple[int, int]:
    idx = y * 12 + (m - 1) + months
    return idx // 12, idx % 12 + 1


def label(y: int, m: int) -> str:
    return f"{MONTH_ABBR[m - 1]}-{str(y)[-2:]}"


def key(y: int, m: int) -> str:
    return f"{y:04d}-{m:02d}"


def delivered_in(db: Session, y: int, m: int) -> list[Engagement]:
    start, end = date(y, m, 1), date(y, m, monthrange(y, m)[1])
    return list(db.scalars(select(Engagement).where(Engagement.actual_date >= start, Engagement.actual_date <= end)
                           .order_by(Engagement.actual_date, Engagement.project_name)))


def available_months(db: Session) -> list[str]:
    dates = db.scalars(select(Engagement.actual_date).where(Engagement.actual_date.is_not(None))).all()
    return sorted({key(d.year, d.month) for d in dates}, reverse=True)


def movement(cur: float, prev: float) -> dict:
    if prev == 0:
        return {"pct": None, "direction": "up" if cur > 0 else "flat",
                "text": "New this period" if cur > 0 else "No change"}
    pct = (cur - prev) / prev * 100
    rounded = int(math.floor(abs(pct) + 0.5))
    if rounded == 0:
        return {"pct": 0, "direction": "flat", "text": "No change"}
    direction = "up" if pct > 0 else "down"
    return {"pct": rounded if pct > 0 else -rounded, "direction": direction,
            "text": f"{'Up' if pct > 0 else 'Down'} by {rounded}%"}


def _by_type(rows, value) -> dict:
    out = {t: 0 for t in PROJECT_TYPES}
    for r in rows:
        out[r.project_type] = out.get(r.project_type, 0) + value(r)
    return out


def _num(v) -> float:
    return v or 0


def _clean(v):
    return int(v) if isinstance(v, float) and v.is_integer() else v


def build_summary(db: Session, month: str | None = None, compare: str = "py") -> dict:
    months = available_months(db)
    if month is None:
        month = months[0] if months else key(date.today().year, date.today().month)
    y, m = parse_month(month)
    cy, cm = shift(y, m, -12) if compare == "py" else shift(y, m, -1)
    cur, prev = delivered_in(db, y, m), delivered_in(db, cy, cm)

    types = list(PROJECT_TYPES)
    if any(r.project_type == "Other" for r in cur + prev):
        types.append("Other")

    def metric(name, value):
        a, b = _by_type(cur, value), _by_type(prev, value)
        a_tot, b_tot = sum(a.values()), sum(b.values())
        return {"name": name,
                "current": {**{t: _clean(a.get(t, 0)) for t in types}, "Total": _clean(a_tot)},
                "compare": {**{t: _clean(b.get(t, 0)) for t in types}, "Total": _clean(b_tot)},
                "movement": movement(a_tot, b_tot)}

    headline = [metric("Number of Engagements Delivered", lambda r: 1),
                metric("Number of KPIs Executed", lambda r: _num(r.kpis_delivered))]

    people: dict[str, Counter] = {}
    for r in cur:
        for role, cell in (("lead", r.delivery_lead), ("primary", r.primary_resource),
                           ("secondary", r.secondary_resource)):
            for name in split_names(cell):
                people.setdefault(name, Counter())[role] += 1
    resources = [{"name": n, "lead": c["lead"], "primary": c["primary"], "secondary": c["secondary"]}
                 for n, c in sorted(people.items())]
    resource_totals = {k: sum(p[k] for p in resources) for k in ("lead", "primary", "secondary")}

    type_rows = [{"type": t, "projects": sum(1 for r in cur if r.project_type == t),
                  "kpis": _clean(sum(_num(r.kpis_delivered) for r in cur if r.project_type == t))} for t in types]

    partners = Counter((r.partner or NOT_SPECIFIED) for r in cur)
    partner_rows = [{"partner": p, "engagements": n}
                    for p, n in sorted(partners.items(), key=lambda kv: (kv[0] == NOT_SPECIFIED, -kv[1], kv[0]))]

    regions = Counter((r.region or NOT_SPECIFIED) for r in cur)
    order = STANDARD_REGIONS + sorted(k for k in regions if k not in STANDARD_REGIONS and k != NOT_SPECIFIED)
    if regions.get(NOT_SPECIFIED):
        order.append(NOT_SPECIFIED)
    region_rows = [{"region": k, "engagements": regions.get(k, 0)} for k in order]

    trend = []
    for i in range(13, -1, -1):  # 14 months ending at the selected month
        ty, tm = shift(y, m, -i)
        rows = delivered_in(db, ty, tm)
        trend.append({"month": key(ty, tm), "label": label(ty, tm), "engagements": len(rows),
                      "kpis": _clean(sum(_num(r.kpis_delivered) for r in rows))})

    last_import = db.scalar(select(ProcessingJob).where(ProcessingJob.job_type == "tracker_import",
                                                        ProcessingJob.status == "completed")
                            .order_by(ProcessingJob.id.desc()))
    missing = (last_import.result or {}).get("missing", []) if last_import else []
    mapping = (last_import.result or {}).get("mapping", {}) if last_import else {}
    unmapped = [FIELDS[f][0] for f in ("partner", "region", "delivery_lead", "primary_resource",
                                        "secondary_resource") if last_import and mapping.get(f) is None]
    return {
        "month": key(y, m), "month_label": label(y, m),
        "compare": compare, "compare_month": key(cy, cm),
        "compare_label": f"{label(cy, cm)} ({'Same month – PY' if compare == 'py' else 'Previous month'})",
        "types": types, "headline": headline,
        "resources": resources, "resource_totals": resource_totals,
        "project_types": type_rows,
        "project_type_totals": {"projects": sum(t["projects"] for t in type_rows),
                                "kpis": _clean(sum(t["kpis"] for t in type_rows))},
        "partners": partner_rows, "regions": region_rows, "trend": trend,
        "months": months,
        "dataset": {"engagements": db.scalar(select(func.count()).select_from(Engagement)) or 0,
                    "not_delivered": db.scalar(select(func.count()).select_from(Engagement)
                                               .where(Engagement.actual_date.is_(None))) or 0,
                    "last_import": ({"job_id": last_import.id, "file": (last_import.result or {}).get("filename"),
                                     "at": last_import.finished_at.isoformat() if last_import.finished_at else None,
                                     "by": last_import.requested_by} if last_import else None),
                    "missing_required": [FIELDS[f][0] for f in missing],
                    "unmapped_optional": unmapped},
    }


def engagement_rows(db: Session, month: str, project_type: str | None = None) -> list[dict]:
    y, m = parse_month(month)
    rows = delivered_in(db, y, m)
    if project_type:
        rows = [r for r in rows if r.project_type == project_type]
    return [{"project_name": r.project_name, "entity_name": r.entity_name, "project_type": r.project_type,
             "kpis_delivered": _clean(r.kpis_delivered), "actual_date": r.actual_date.isoformat(),
             "partner": r.partner, "region": r.region, "delivery_lead": r.delivery_lead,
             "primary_resource": r.primary_resource, "secondary_resource": r.secondary_resource,
             "business_area": r.business_area} for r in rows]


# ---------------------------------------------------------------- Excel export
PEACH, BLUE, PURPLE, GREY = "F4B183", "9DD9F3", "7030A0", "D9D9D9"
PEACH_L, BLUE_L = "FBE5D6", "DDEBF7"


def build_excel(summary: dict, engagements: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "MIS Summary"
    thin = Side(style="thin", color="000000")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def put(r, c, v=None, fill=None, font=None, align=None, fmt=None):
        cell = ws.cell(r, c, v)
        cell.border = box
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        if font:
            cell.font = font
        cell.alignment = align or Alignment(vertical="center", wrap_text=True)
        if fmt:
            cell.number_format = fmt
        return cell

    types = summary["types"]
    n = len(types) + 1
    ws.cell(1, 1, f"MIS summary – {summary['month_label']}").font = Font(bold=True, size=14)

    # headline table
    r0 = 3
    put(r0, 1, "Description", font=bold, align=center)
    ws.merge_cells(start_row=r0, start_column=1, end_row=r0 + 1, end_column=1)
    put(r0, 2, summary["month_label"], PEACH, bold, center)
    ws.merge_cells(start_row=r0, start_column=2, end_row=r0, end_column=1 + n)
    put(r0, 2 + n, summary["compare_label"], BLUE, bold, center)
    ws.merge_cells(start_row=r0, start_column=2 + n, end_row=r0, end_column=1 + 2 * n)
    mv = 2 + 2 * n
    put(r0, mv, "Movement (in %)", font=bold, align=center)
    ws.merge_cells(start_row=r0, start_column=mv, end_row=r0 + 1, end_column=mv)
    for i, t in enumerate(types + ["Total"]):
        put(r0 + 1, 2 + i, t, PEACH_L, bold, center)
        put(r0 + 1, 2 + n + i, t, BLUE_L, bold, center)
    for k, h in enumerate(summary["headline"]):
        r = r0 + 2 + k
        put(r, 1, h["name"])
        for i, t in enumerate(types + ["Total"]):
            put(r, 2 + i, h["current"][t], font=bold if t == "Total" else None, align=center)
            put(r, 2 + n + i, h["compare"][t], font=bold if t == "Total" else None, align=center)
        color = {"up": "548235", "down": "C00000"}.get(h["movement"]["direction"], "595959")
        put(r, mv, h["movement"]["text"], font=Font(color=color), align=center)
        ws.row_dimensions[r].height = 32
    ws.column_dimensions["A"].width = 30
    for c in range(2, mv + 1):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = 11
    ws.column_dimensions[ws.cell(1, mv).column_letter].width = 18

    def table(r, headers, rows, total=None, widths=None):
        for i, h in enumerate(headers):
            put(r, 1 + i, h, PURPLE, Font(bold=True, color="FFFFFF"))
        for k, row in enumerate(rows, start=1):
            for i, v in enumerate(row):
                put(r + k, 1 + i, v, align=Alignment(horizontal="left" if i == 0 else "right"))
        end = r + len(rows)
        if total:
            end += 1
            for i, v in enumerate(total):
                put(end, 1 + i, v, GREY, bold, Alignment(horizontal="left" if i == 0 else "right"))
        return end + 2

    r = r0 + 5
    r = table(r, ["Name", "As Delivery Lead", "As Primary Resource", "As Secondary Resource"],
              [[p["name"], p["lead"], p["primary"], p["secondary"]] for p in summary["resources"]],
              ["Total", summary["resource_totals"]["lead"], summary["resource_totals"]["primary"],
               summary["resource_totals"]["secondary"]])
    r = table(r, ["Project Type", "No. of Projects", "No. of KPIs executed"],
              [[t["type"], t["projects"], t["kpis"]] for t in summary["project_types"]],
              ["Total", summary["project_type_totals"]["projects"], summary["project_type_totals"]["kpis"]])
    table(r, ["Partner Name", "No. of Engagements"], [[p["partner"], p["engagements"]] for p in summary["partners"]],
          ["Total", sum(p["engagements"] for p in summary["partners"])])

    # chart data + charts
    cd = wb.create_sheet("Chart data")
    cd.append(["Region", "Engagements"])
    for x in summary["regions"]:
        cd.append([x["region"], x["engagements"]])
    start = len(summary["regions"]) + 3
    cd.cell(start, 1, "Month")
    cd.cell(start, 2, "Engagements delivered")
    cd.cell(start, 3, "KPIs executed")
    for i, t in enumerate(summary["trend"], start=1):
        cd.cell(start + i, 1, t["label"])
        cd.cell(start + i, 2, t["engagements"])
        cd.cell(start + i, 3, t["kpis"])

    pie = PieChart()
    pie.title = "Engagement Distribution by Region"
    pie.add_data(Reference(cd, min_col=2, min_row=1, max_row=1 + len(summary["regions"])), titles_from_data=True)
    pie.set_categories(Reference(cd, min_col=1, min_row=2, max_row=1 + len(summary["regions"])))
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showVal = True
    pie.height, pie.width = 8, 12
    anchor_col = ws.cell(1, mv + 2).column_letter
    ws.add_chart(pie, f"{anchor_col}3")

    line = LineChart()
    line.title = "Month by month – engagements delivered"
    line.add_data(Reference(cd, min_col=2, min_row=start, max_row=start + len(summary["trend"])),
                  titles_from_data=True)
    line.set_categories(Reference(cd, min_col=1, min_row=start + 1, max_row=start + len(summary["trend"])))
    line.legend = None
    line.height, line.width = 8, 20
    ws.add_chart(line, f"{anchor_col}21")

    det = wb.create_sheet("Delivered engagements")
    cols = [("project_name", "Project Name"), ("entity_name", "Entity Name"), ("project_type", "Project Type"),
            ("kpis_delivered", "KPIs Delivered"), ("actual_date", "Actual Date"), ("partner", "Partner"),
            ("region", "Region"), ("delivery_lead", "Delivery Lead"), ("primary_resource", "Primary Resource"),
            ("secondary_resource", "Secondary Resource"), ("business_area", "Business Area")]
    det.append([c[1] for c in cols])
    for c in det[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=PURPLE)
    for e in engagements:
        det.append([e[k] for k, _ in cols])
    for i, w in enumerate([28, 34, 12, 14, 12, 20, 10, 18, 18, 18, 18], start=1):
        det.column_dimensions[det.cell(1, i).column_letter].width = w
    det.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
