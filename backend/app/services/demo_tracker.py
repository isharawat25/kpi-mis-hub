"""Deterministic demo engagements and a sample tracker workbook laid out like the team's real file
(legend block, two-row header with a merged 'Step 4 - KPI Execution' group, blank cells).
All names are invented."""
import io
import random
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

LEADS = ["Kiran Rao", "Meera Nair", "Rohit Sen"]
RESOURCES = ["Aarav Jain", "Bhavna Kulkarni", "Chirag Shah", "Deepa Menon", "Esha Verma",
             "Farhan Ali", "Gauri Patil", "Harsh Malhotra"]
PARTNERS = ["Arjun Mehra", "Neha Bhat", "Suresh Pillai", "Ritu Ahuja", "Vikram Desai",
            "Pooja Sinha", "Manoj Thakur", "Kavita Rao", "Sanjay Gupta", "Lata Krishnan"]
COMPANIES = ["Alder", "Birch", "Cedar", "Delta", "Ember", "Fjord", "Garnet", "Harbor", "Indigo", "Juniper",
             "Kestrel", "Lumen", "Maple", "Nimbus", "Onyx", "Pinnacle", "Quartz", "Ridge", "Sable", "Topaz"]
SUFFIX = ["Private Limited", "Limited", "India Private Limited", "Technologies Private Limited"]
REGIONS = (["North"] * 9 + ["South"] * 3 + ["East"] * 1 + ["West"] * 7)
AREAS = ["General Ledger", "General Ledger", "General Ledger", "Revenue", "Procure to Pay", "Payroll"]
ERPS = ["SAP", "SAP", "Oracle", "Tally", "Microsoft Dynamics", None]

# engagements delivered per month, Sep-2025 .. Oct-2026
MONTHS = [(2025, 9), (2025, 10), (2025, 11), (2025, 12), (2026, 1), (2026, 2), (2026, 3), (2026, 4),
          (2026, 5), (2026, 6), (2026, 7), (2026, 8), (2026, 9), (2026, 10)]
VOLUMES = [22, 20, 14, 12, 18, 26, 21, 48, 55, 40, 34, 30, 44, 25]


def demo_records() -> list[dict]:
    rng = random.Random(42)
    out = []
    n = 0
    for (y, m), vol in zip(MONTHS, VOLUMES):
        for _ in range(vol):
            n += 1
            out.append(_record(rng, n, date(y, m, rng.randint(1, 28))))
    for _ in range(18):  # open engagements, not yet delivered
        n += 1
        rec = _record(rng, n, None)
        rec["expected_date"] = date(2026, rng.choice([10, 11, 12]), rng.randint(1, 28))
        rec["kpis_delivered"] = 0
        out.append(rec)
    return out


def _record(rng: random.Random, n: int, actual: date | None) -> dict:
    roll = rng.random()
    ptype = "Stat" if roll < 0.84 else ("IA" if roll < 0.95 else "Tax")
    kpis = {"Stat": rng.randint(5, 22), "IA": rng.randint(18, 45), "Tax": rng.randint(8, 30)}[ptype]
    company = f"{rng.choice(COMPANIES)} {rng.choice(COMPANIES)}"
    lead = rng.choice(LEADS) if rng.random() < 0.97 else None
    primary = rng.choice(RESOURCES)
    secondary = rng.choice([r for r in RESOURCES if r != primary]) if rng.random() < 0.16 else None
    return {
        "project_name": f"{company} {n:03d}",
        "entity_name": f"{company} {rng.choice(SUFFIX)}",
        "project_type": ptype,
        "kpis_planned": kpis,
        "kpis_delivered": kpis if actual else None,
        "expected_date": (date(actual.year, actual.month, min(28, actual.day + rng.randint(0, 5)))
                          if actual else None),
        "actual_date": actual,
        "business_area": rng.choice(AREAS),
        "engagement_period": "Apr25-Mar26" if (actual is None or actual < date(2026, 4, 1)) else "Apr26-Mar27",
        "erp": rng.choice(ERPS),
        "partner": rng.choice(PARTNERS),
        "region": rng.choice(REGIONS),
        "delivery_lead": lead,
        "primary_resource": primary,
        "secondary_resource": secondary,
        "file_count": rng.choice([None, 1, 2, 3, 4]),
        "file_size_mb": rng.choice([None, round(rng.uniform(5, 3000), 1)]),
    }


COLUMNS = [  # (top header, sub header, record key)
    ("Sno", None, None), ("Project Name", None, "project_name"), ("Entity Name", None, "entity_name"),
    ("Project Type", None, "project_type"), ("Delivery Lead", None, "delivery_lead"),
    ("Primary Resource", None, "primary_resource"), ("Secondary Resource", None, "secondary_resource"),
    ("Partner Name", None, "partner"), ("Region", None, "region"),
    ("Step 4 - KPI Execution", "Number of KPIs in scope", "kpis_planned"),
    ("Step 4 - KPI Execution", "Number of KPIs Delivered", "kpis_delivered"),
    ("Expected Date of Delivery", None, "expected_date"), ("Actual Date of Delivery", None, "actual_date"),
    ("Business area", None, "business_area"), ("Engagement period (year / month)", None, "engagement_period"),
    ("ERP", None, "erp"), ("File size (in MB)", None, "file_size_mb"), ("No of Files", None, "file_count"),
]


def build_sample_tracker(records: list[dict] | None = None) -> bytes:
    records = records if records is not None else demo_records()
    wb = Workbook()
    ws = wb.active
    ws.title = "Tracker"
    thin = Side(style="thin", color="999999")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    legend = [("FF0000", "Overdue"), ("FFFF00", "Due this month"), ("92D050", "Delivered"),
              ("FFFFFF", "Not started"), ("D9D9D9", "On hold"), ("8EA9DB", "In progress"), ("808080", "Dropped")]
    for i, (color, label) in enumerate(legend, start=2):
        ws.cell(i, 2).fill = PatternFill("solid", fgColor=color)
        ws.cell(i, 2).border = border
        ws.cell(i, 3, label).border = border
    hdr = 13
    gold = PatternFill("solid", fgColor="FFC000")
    for j, (top, sub, _) in enumerate(COLUMNS, start=1):
        ws.cell(hdr, j, top)
        if sub:
            ws.cell(hdr + 1, j, sub)
        for r in (hdr, hdr + 1):
            c = ws.cell(r, j)
            c.fill, c.border = gold, border
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    j = 1
    while j <= len(COLUMNS):  # merge vertical headers and the grouped Step 4 header
        top, sub, _ = COLUMNS[j - 1]
        if sub is None:
            ws.merge_cells(start_row=hdr, start_column=j, end_row=hdr + 1, end_column=j)
            j += 1
        else:
            k = j
            while k < len(COLUMNS) and COLUMNS[k][0] == top:
                k += 1
            ws.merge_cells(start_row=hdr, start_column=j, end_row=hdr, end_column=k)
            j = k + 1
    for i, rec in enumerate(records, start=1):
        r = hdr + 1 + i
        for j, (_, _, key) in enumerate(COLUMNS, start=1):
            value = i if key is None else rec.get(key)
            c = ws.cell(r, j, value)
            c.border = border
            if isinstance(value, date):
                c.number_format = "yyyy-mm-dd"
    widths = [6, 26, 34, 10, 16, 18, 18, 18, 9, 12, 12, 14, 14, 18, 16, 16, 12, 10]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(hdr + 2, 3)
    ws.auto_filter.ref = f"A{hdr + 1}:{get_column_letter(len(COLUMNS))}{hdr + 1 + len(records)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
