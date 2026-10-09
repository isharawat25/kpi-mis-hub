"""Phase 2: validate and import Excel/CSV KPI submissions."""
import io
import math
from datetime import datetime, timezone

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import KPI, BusinessUnit, Document, KPIValue, ProcessingJob, RejectedRecord, ReportingPeriod
from .logs import log
from .storage import get_storage

REQUIRED_COLUMNS = ["kpi_code", "business_unit", "period", "actual"]
OPTIONAL_COLUMNS = ["target", "commentary"]
TEMPLATE_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS


class TemplateError(ValueError):
    pass


def _norm(col: str) -> str:
    return str(col).strip().lower().replace(" ", "_").replace("-", "_")


def read_submission(data: bytes, filename: str) -> pd.DataFrame:
    name = filename.lower()
    try:
        if name.endswith(".csv"):
            df = pd.read_csv(io.BytesIO(data), dtype=str, keep_default_na=False)
        elif name.endswith((".xlsx", ".xlsm")):
            xls = pd.ExcelFile(io.BytesIO(data), engine="openpyxl")
            sheet = "Submission" if "Submission" in xls.sheet_names else xls.sheet_names[0]
            df = pd.read_excel(xls, sheet_name=sheet, dtype=str, keep_default_na=False)
        else:
            raise TemplateError("Unsupported file type. Upload .xlsx or .csv")
    except TemplateError:
        raise
    except Exception as e:  # corrupt file etc.
        raise TemplateError(f"Could not read file: {e}") from e

    df.columns = [_norm(c) for c in df.columns]
    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise TemplateError(f"Template mismatch - missing columns: {', '.join(missing)}")
    for c in OPTIONAL_COLUMNS:
        if c not in df.columns:
            df[c] = ""
    df = df[TEMPLATE_COLUMNS]
    df = df[~(df.apply(lambda r: all(str(v).strip() == "" for v in r), axis=1))]  # drop blank rows
    if df.empty:
        raise TemplateError("File contains no data rows")
    return df


def parse_number(value) -> float | None:
    """Accepts '1,234.5', '(100)' for negatives, ' 12 '. Returns None for blank. Raises on garbage."""
    s = str(value).strip()
    if s == "":
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace(",", "").replace(" ", "")
    num = float(s)  # ValueError propagates
    if math.isnan(num) or math.isinf(num):
        raise ValueError("not a finite number")
    return -num if neg else num


def process_submission_job(db: Session, job_id: int) -> ProcessingJob:
    job = db.get(ProcessingJob, job_id)
    doc = db.get(Document, job.document_id)
    job.status, job.started_at = "running", datetime.now(timezone.utc)
    log(db, job.id, f"Processing file '{doc.filename}'")
    db.commit()

    try:
        df = read_submission(get_storage().get(doc.blob_path), doc.filename)
    except TemplateError as e:
        job.status, job.error, job.finished_at = "failed", str(e), datetime.now(timezone.utc)
        log(db, job.id, str(e), "ERROR")
        db.commit()
        return job

    kpis = {k.code.upper(): k for k in db.scalars(select(KPI))}
    bus = {b.code.upper(): b for b in db.scalars(select(BusinessUnit))}
    periods = {p.code.upper(): p for p in db.scalars(select(ReportingPeriod))}

    seen, accepted, rejected = set(), 0, 0
    now = datetime.now(timezone.utc)
    for idx, row in df.iterrows():
        row_no = int(idx) + 2  # header is row 1
        raw = {k: str(v) for k, v in row.items()}
        errors = []
        kcode, bcode, pcode = (str(row[c]).strip().upper() for c in ("kpi_code", "business_unit", "period"))

        kpi, bu, period = kpis.get(kcode), bus.get(bcode), periods.get(pcode)
        if not kcode:
            errors.append("kpi_code is required")
        elif not kpi or not kpi.is_active:
            errors.append(f"Unknown or inactive KPI '{kcode}'")
        elif kpi.kpi_type != "input":
            errors.append(f"KPI '{kcode}' is calculated by formula and cannot be submitted")
        if not bcode:
            errors.append("business_unit is required")
        elif not bu or not bu.is_active:
            errors.append(f"Unknown or inactive business unit '{bcode}'")
        if not pcode:
            errors.append("period is required")
        elif not period:
            errors.append(f"Unknown reporting period '{pcode}'")
        elif period.status != "open":
            errors.append(f"Period '{pcode}' is closed")
        elif kpi and kpi.frequency != period.frequency:
            errors.append(f"KPI '{kcode}' is {kpi.frequency} but period '{pcode}' is {period.frequency}")

        actual = target = None
        try:
            actual = parse_number(row["actual"])
            if actual is None:
                errors.append("actual is required")
        except ValueError:
            errors.append(f"actual '{row['actual']}' is not a valid number")
        try:
            target = parse_number(row["target"])
        except ValueError:
            errors.append(f"target '{row['target']}' is not a valid number")

        key = (kcode, bcode, pcode)
        if key in seen:
            errors.append("Duplicate KPI/business unit/period row in this file")
        seen.add(key)

        existing = None
        if not errors:
            existing = db.scalar(select(KPIValue).where(
                KPIValue.kpi_id == kpi.id, KPIValue.business_unit_id == bu.id, KPIValue.period_id == period.id))
            if existing and existing.status == "approved":
                errors.append("Value already approved for this KPI/unit/period; ask an approver to reopen it")

        if errors:
            rejected += 1
            db.add(RejectedRecord(job_id=job.id, row_number=row_no, raw_data=raw, reasons=errors))
            continue

        v = existing or KPIValue(kpi_id=kpi.id, business_unit_id=bu.id, period_id=period.id)
        v.actual = actual
        if target is not None:
            v.target = target
        if str(row["commentary"]).strip():
            v.commentary = str(row["commentary"]).strip()
        v.status, v.source, v.source_job_id = "submitted", "file", job.id
        v.submitted_by, v.submitted_at = job.requested_by, now
        v.reviewed_by = v.reviewed_at = v.review_note = None
        if not existing:
            db.add(v)
        accepted += 1

    job.rows_total, job.rows_accepted, job.rows_rejected = len(df), accepted, rejected
    job.status, job.finished_at = "completed", datetime.now(timezone.utc)
    log(db, job.id, f"Completed: {accepted} accepted, {rejected} rejected of {len(df)} rows",
        "WARN" if rejected else "INFO")
    db.commit()
    return job


def build_template(db: Session) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Submission"
    ws.append(TEMPLATE_COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    kpis = db.scalars(select(KPI).where(KPI.kpi_type == "input", KPI.is_active)).all()
    bus = db.scalars(select(BusinessUnit).where(BusinessUnit.is_active)).all()
    period = db.scalar(select(ReportingPeriod).where(ReportingPeriod.status == "open")
                       .order_by(ReportingPeriod.start_date.desc()))
    if kpis and bus and period:
        for k in kpis:
            if k.frequency == period.frequency:
                ws.append([k.code, bus[0].code, period.code, None, None, ""])
    for col, w in zip("ABCDEF", (18, 16, 12, 14, 14, 50)):
        ws.column_dimensions[col].width = w

    ref = wb.create_sheet("Reference")
    ref.append(["Type", "Code", "Name", "Extra"])
    for k in kpis:
        ref.append(["KPI", k.code, k.name, f"{k.frequency}, unit={k.unit}"])
    for b in bus:
        ref.append(["Business unit", b.code, b.name, ""])
    for p in db.scalars(select(ReportingPeriod).where(ReportingPeriod.status == "open")):
        ref.append(["Period", p.code, p.label, p.frequency])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rejected_report(db: Session, job_id: int) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rejected records"
    ws.append(["row_number", *TEMPLATE_COLUMNS, "reasons"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in db.scalars(select(RejectedRecord).where(RejectedRecord.job_id == job_id)
                        .order_by(RejectedRecord.row_number)):
        ws.append([r.row_number, *[r.raw_data.get(c, "") for c in TEMPLATE_COLUMNS], "; ".join(r.reasons)])
    ws.column_dimensions["H"].width = 80
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
