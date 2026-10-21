"""Import the team's project tracker (one row per engagement) exactly as it is kept in Excel.

Handles legend/notes rows above the table, two-row headers with merged parent cells
(e.g. "Step 4 - KPI Execution" over "Number of KPIs Delivered"), blank cells and mixed date formats.
Columns are matched by header words; users can override the match and it is remembered.
"""
import io
import re
from datetime import date, datetime, timedelta, timezone

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from ..models import AppSetting, Document, Engagement, ProcessingJob
from .logs import log
from .storage import get_storage

MAPPING_SETTING = "tracker_mapping"

# field -> (label, kind, required, [(keywords, score)], excluded words)
FIELDS: dict[str, tuple] = {
    "project_name": ("Project name", "text", True,
                     [(["project name"], 100), (["engagement name"], 90), (["client name"], 60)], []),
    "entity_name": ("Entity name", "text", False, [(["entity"], 100)], []),
    "project_type": ("Project type (Stat / IA / Tax)", "text", True,
                     [(["project type"], 100), (["engagement type"], 90), (["type"], 40)], ["erp", "file"]),
    "kpis_planned": ("Number of KPIs (planned / in scope)", "number", False,
                     [(["kpi", "plan"], 100), (["kpi", "scope"], 95), (["kpi", "total"], 80),
                      (["kpi", "execution"], 70), (["number of kpi"], 60)],
                     ["deliver", "business", "executed", "date"]),
    "kpis_delivered": ("Number of KPIs delivered", "number", True,
                       [(["kpi", "deliver"], 100), (["kpi", "executed"], 90), (["kpi", "complete"], 80)],
                       ["business", "date"]),
    "expected_date": ("Expected date of delivery", "date", False,
                      [(["expected", "date"], 100), (["expected"], 80), (["due", "date"], 80),
                       (["target", "date"], 70)], ["actual"]),
    "actual_date": ("Actual date of delivery", "date", True,
                    [(["actual", "deliver"], 110), (["actual", "date"], 100), (["delivery", "date"], 90),
                     (["completion", "date"], 90), (["actual"], 70)], ["expected"]),
    "business_area": ("Business area", "text", False, [(["business area"], 100)], ["kpi", "count"]),
    "engagement_period": ("Engagement period", "text", False,
                          [(["engagement period"], 100), (["period"], 50)], []),
    "erp": ("ERP", "text", False, [(["erp"], 100)], []),
    "partner": ("Partner name", "text", False,
                [(["engagement partner"], 110), (["partner"], 100)], []),
    "region": ("Region", "text", False, [(["region"], 100), (["zone"], 70), (["location"], 50)], []),
    "delivery_lead": ("Delivery lead", "text", False,
                      [(["delivery lead"], 110), (["lead"], 70)], ["primary", "secondary"]),
    "primary_resource": ("Primary resource", "text", False, [(["primary"], 100)], []),
    "secondary_resource": ("Secondary resource", "text", False, [(["secondary"], 100)], []),
    "file_count": ("No. of files", "number", False,
                   [(["no of files"], 100), (["number of files"], 100), (["files"], 60)], ["size"]),
    "file_size_mb": ("File size (MB)", "number", False, [(["file size"], 100), (["size"], 60)], []),
}

PROJECT_TYPES = ["Stat", "IA", "Tax"]


class TrackerError(ValueError):
    pass


def _norm(text) -> str:
    s = re.sub(r"[^a-z0-9]+", " ", str(text or "").lower())
    return re.sub(r"\s+", " ", s).strip()


def _has(header: str, keyword: str) -> bool:
    """Phrase keywords match at a word start; single words match token prefixes (deliver -> delivered)."""
    if " " in keyword:
        return re.search(r"(^| )" + re.escape(keyword), header) is not None
    return any(tok.startswith(keyword) for tok in header.split())


def _score(field: str, header: str) -> int:
    _, _, _, patterns, excluded = FIELDS[field]
    if not header or any(_has(header, x) for x in excluded):
        return 0
    return max((s for words, s in patterns if all(_has(header, w) for w in words)), default=0)


# ---------------------------------------------------------------- reading the sheet
def _cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_table(data: bytes, filename: str) -> dict:
    """Returns {"sheet", "header_row", "headers": [{index, letter, text}], "rows": [(excel_row, [values])]}."""
    name = filename.lower()
    if name.endswith(".csv"):
        df = pd.read_csv(io.BytesIO(data), header=None, dtype=object, keep_default_na=False)
        grid = [[None if (isinstance(v, str) and v.strip() == "") else v for v in row] for row in df.values.tolist()]
        return _locate(grid, sheet="CSV")
    if not name.endswith((".xlsx", ".xlsm")):
        raise TrackerError("Upload the tracker as .xlsx or .csv")
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:  # noqa: BLE001
        raise TrackerError(f"Could not open the workbook: {e}") from e
    sheets = [wb.active] + [ws for ws in wb.worksheets if ws is not wb.active]
    last_error = None
    for ws in sheets:
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        for rng in ws.merged_cells.ranges:  # spread merged header values to every covered cell
            if rng.min_row > 60:
                continue
            top = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, min(rng.max_row, len(grid)) + 1):
                row = grid[r - 1]
                for c in range(rng.min_col, rng.max_col + 1):
                    if c - 1 < len(row):
                        row[c - 1] = top
        try:
            return _locate(grid, sheet=ws.title)
        except TrackerError as e:
            last_error = e
    raise last_error or TrackerError("No sheet with a 'Project Name' column was found")


def _locate(grid: list[list], sheet: str) -> dict:
    header_idx = name_col = None
    for i, row in enumerate(grid[:60]):
        for j, v in enumerate(row):
            if _score("project_name", _norm(v)) >= 90:
                header_idx, name_col = i, j
                break
        if header_idx is not None:
            break
    if header_idx is None:
        raise TrackerError("Could not find the header row: no column called 'Project Name' in the first 60 rows")

    width = max(len(r) for r in grid[header_idx: header_idx + 2])
    top = grid[header_idx] + [None] * (width - len(grid[header_idx]))
    sub = grid[header_idx + 1] + [None] * (width - len(grid[header_idx + 1])) if header_idx + 1 < len(grid) else []
    # second header row: nothing (or the same merged text) under "Project Name", and some text elsewhere
    is_sub = bool(sub) and (_cell_text(sub[name_col]) in ("", _cell_text(top[name_col]))) and any(
        isinstance(v, str) and v.strip() for v in sub)
    headers = []
    for j in range(width):
        parts = [_cell_text(top[j])]
        if is_sub and _cell_text(sub[j]) and _cell_text(sub[j]) != parts[0]:
            parts.append(_cell_text(sub[j]))
        text = " ".join(p for p in parts if p)
        headers.append({"index": j, "letter": get_column_letter(j + 1), "text": text})
    first = header_idx + (2 if is_sub else 1)
    rows = []
    for i in range(first, len(grid)):
        row = list(grid[i]) + [None] * (width - len(grid[i]))
        rows.append((i + 1, row[:width]))
    return {"sheet": sheet, "header_row": header_idx + 1, "headers": headers, "rows": rows}


# ---------------------------------------------------------------- column matching
def match_columns(headers: list[dict], saved: dict | None = None, overrides: dict | None = None) -> dict:
    """field -> column index (or None). Overrides (index or None) win, then saved header texts, then scoring."""
    mapping: dict[str, int | None] = {f: None for f in FIELDS}
    used: set[int] = set()
    by_norm = {_norm(h["text"]): h["index"] for h in headers if h["text"]}
    for f, idx in (overrides or {}).items():
        if f in FIELDS:
            mapping[f] = None if idx in (None, "", -1) else int(idx)
            if mapping[f] is not None:
                used.add(mapping[f])
    fixed = set((overrides or {}).keys())
    for f, text in (saved or {}).items():
        if f in FIELDS and f not in fixed and text and _norm(text) in by_norm and by_norm[_norm(text)] not in used:
            mapping[f] = by_norm[_norm(text)]
            used.add(mapping[f])
            fixed.add(f)
    candidates = sorted(((_score(f, _norm(h["text"])), f, h["index"]) for f in FIELDS if f not in fixed
                         for h in headers), key=lambda t: (-t[0], t[2]))
    for score, f, idx in candidates:
        if score > 0 and mapping[f] is None and f not in fixed and idx not in used:
            mapping[f] = idx
            used.add(idx)
    return mapping


# ---------------------------------------------------------------- value parsing
def _text(v) -> str | None:
    s = _cell_text(v)
    s = re.sub(r"\s+", " ", s)
    return None if s.lower() in ("", "-", "na", "n/a", "nil", "none") else s


def _number(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = _text(v)
    if s is None:
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace(",", "").replace(" ", "")
    num = float(s)  # ValueError -> caller warns
    return -num if neg else num


def _date(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 20000 < v < 80000:  # Excel serial day number
            return date(1899, 12, 30) + timedelta(days=int(v))
        raise ValueError("number is not a date")
    s = _text(v)
    if s is None:
        return None
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}([ T].*)?", s):
        return date.fromisoformat(s[:10])
    parsed = pd.to_datetime(s, dayfirst=True, errors="raise")
    return parsed.date()


def normalise_type(raw: str | None) -> str:
    s = _norm(raw)
    if not s:
        return "Other"
    if s.startswith("stat"):
        return "Stat"
    if s in ("ia", "i a") or s.startswith("internal"):
        return "IA"
    if s.startswith("tax"):
        return "Tax"
    return "Other"


def normalise_region(raw: str | None) -> str | None:
    s = _text(raw)
    return s.title() if s else None


def split_names(cell: str | None) -> list[str]:
    if not cell:
        return []
    parts = re.split(r"\s*(?:,|/|&|;|\n|\band\b)\s*", cell)
    return [re.sub(r"\s+", " ", p).strip().title() for p in parts if p and p.strip()]


# ---------------------------------------------------------------- import
def run_import(db: Session, job_id: int, overrides: dict | None = None, remember: bool = True) -> ProcessingJob:
    job = db.get(ProcessingJob, job_id)
    doc = db.get(Document, job.document_id)
    params = job.result or {}
    mode = params.get("mode", "replace")
    job.status, job.started_at = "running", datetime.now(timezone.utc)
    db.commit()
    try:
        table = read_table(get_storage().get(doc.blob_path), doc.filename)
    except TrackerError as e:
        job.status, job.error, job.finished_at = "failed", str(e), datetime.now(timezone.utc)
        log(db, job.id, str(e), "ERROR")
        db.commit()
        return job

    saved_row = db.get(AppSetting, MAPPING_SETTING)
    saved = saved_row.value if saved_row else None
    mapping = match_columns(table["headers"], saved, overrides)
    headers = table["headers"]
    missing = [f for f, (_, _, req, _, _) in FIELDS.items() if req and mapping[f] is None]
    if mapping["project_name"] is None:
        job.status, job.finished_at = "failed", datetime.now(timezone.utc)
        job.error = "No 'Project Name' column could be matched"
        job.result = {**params, "headers": headers, "mapping": mapping, "missing": missing, "sheet": table["sheet"]}
        db.commit()
        return job

    warnings: dict[str, dict] = {}

    def warn(kind: str, row_no: int, detail: str):
        w = warnings.setdefault(kind, {"count": 0, "rows": []})
        w["count"] += 1
        if len(w["rows"]) < 5:
            w["rows"].append(f"row {row_no}: {detail}" if detail else f"row {row_no}")

    def get(row, field):
        idx = mapping[field]
        return row[idx] if idx is not None and idx < len(row) else None

    records = []
    for row_no, row in table["rows"]:
        name = _text(get(row, "project_name"))
        if not name or _norm(name) in ("total", "grand total", "sub total", "subtotal"):
            continue
        rec = {"source_row": row_no, "project_name": name}
        for field, (label, kind, _, _, _) in FIELDS.items():
            if field == "project_name" or mapping[field] is None:
                continue
            raw = get(row, field)
            try:
                if kind == "number":
                    rec[field] = _number(raw)
                elif kind == "date":
                    rec[field] = _date(raw)
                else:
                    rec[field] = _text(raw)
            except (ValueError, TypeError, OverflowError):
                rec[field] = None
                warn(f"{label}: not a valid {kind}", row_no, _cell_text(raw))
        raw_type = rec.pop("project_type", None)
        rec["project_type_raw"] = raw_type
        rec["project_type"] = normalise_type(raw_type)
        rec["region"] = normalise_region(rec.get("region"))
        if mapping["project_type"] is not None and rec["project_type"] == "Other":
            warn("Project type is blank or not Stat / IA / Tax (counted as Other)", row_no, raw_type or "blank")
        if rec.get("actual_date") and rec.get("kpis_delivered") is None and mapping["kpis_delivered"] is not None:
            warn("Delivered (has actual date) but number of KPIs delivered is blank", row_no, name)
        records.append(rec)

    if mode == "replace":
        db.execute(delete(Engagement))
    else:  # append: a later import replaces the same project + entity + delivery date
        for rec in records:
            db.execute(delete(Engagement).where(
                Engagement.project_name == rec["project_name"],
                Engagement.entity_name.is_(None) if rec.get("entity_name") is None
                else Engagement.entity_name == rec["entity_name"],
                Engagement.actual_date.is_(None) if rec.get("actual_date") is None
                else Engagement.actual_date == rec["actual_date"]))
    for rec in records:
        db.add(Engagement(import_job_id=job.id, **rec))

    if remember:
        text_map = {f: headers[i]["text"] for f, i in mapping.items() if i is not None}
        if saved_row:
            saved_row.value = text_map
        else:
            db.add(AppSetting(key=MAPPING_SETTING, value=text_map))

    delivered = sum(1 for r in records if r.get("actual_date"))
    job.rows_total = len(records)
    job.rows_accepted = delivered
    job.rows_rejected = 0
    job.result = {**params, "sheet": table["sheet"], "header_row": table["header_row"], "headers": headers,
                  "mapping": mapping, "missing": missing, "warnings": warnings,
                  "engagements": len(records), "delivered": delivered, "not_delivered": len(records) - delivered}
    job.status, job.finished_at, job.error = "completed", datetime.now(timezone.utc), None
    log(db, job.id, f"Tracker imported ({mode}): {len(records)} engagements, {delivered} with a delivery date"
        + (f"; missing columns: {', '.join(FIELDS[m][0] for m in missing)}" if missing else ""),
        "WARN" if missing or warnings else "INFO")
    db.commit()
    return job


def field_list() -> list[dict]:
    return [{"field": f, "label": v[0], "kind": v[1], "required": v[2]} for f, v in FIELDS.items()]
