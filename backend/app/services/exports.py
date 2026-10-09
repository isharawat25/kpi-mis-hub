"""Presentation-ready outputs: formatted Excel workbook and PowerPoint tables."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

RAG_FILL = {"GREEN": "C6EFCE", "AMBER": "FFEB9C", "RED": "FFC7CE", "NA": "EDEDED"}
HEADERS = ["Business Unit", "KPI Code", "KPI", "Unit", "Actual", "Target", "Variance", "Variance %",
           "Prior Actual", "Change vs Prior %", "RAG", "Calculated", "Note"]


def _row(r):
    return [r.business_unit, r.kpi_code, r.kpi_name, r.unit, r.actual, r.target, r.variance, r.variance_pct,
            r.prior_actual, r.change_vs_prior_pct, r.rag, "Yes" if r.is_calculated else "No", r.note or ""]


def _sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    ws.append(HEADERS)
    thin = Side(style="thin", color="BFBFBF")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append(_row(r))
        rag_cell = ws.cell(row=ws.max_row, column=11)
        rag_cell.fill = PatternFill("solid", fgColor=RAG_FILL.get(r.rag, "FFFFFF"))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        for idx in (5, 6, 7, 9):
            row[idx - 1].number_format = "#,##0.00"
        for idx in (8, 10):
            row[idx - 1].number_format = '0.00"%"'
    widths = [14, 16, 32, 8, 14, 14, 14, 12, 14, 14, 8, 10, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def build_excel(period, rows) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Group Summary", [r for r in rows if r.business_unit == "GROUP"])
    _sheet(wb, "By Business Unit", [r for r in rows if r.business_unit != "GROUP"])
    _sheet(wb, "Exceptions", [r for r in rows if r.rag in ("RED", "AMBER")])
    _sheet(wb, "Consolidated Dataset", rows)
    info = wb.create_sheet("About", 0)
    info["A1"] = f"MIS Pack - {period.label} ({period.code})"
    info["A1"].font = Font(bold=True, size=14)
    info["A3"] = "All figures are calculated deterministically from submitted KPI values and approved formulas."
    info["A4"] = "RAG: GREEN = on/above target, AMBER = within tolerance, RED = outside tolerance."
    info.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fmt(v, pct=False):
    if v is None:
        return "-"
    return f"{v:,.2f}%" if pct else f"{v:,.2f}"


def _table_slide(prs, title, rows):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(26)
    cols = ["KPI", "Actual", "Target", "Var %", "vs Prior %", "RAG"]
    table = slide.shapes.add_table(len(rows) + 1, len(cols), Inches(0.4), Inches(1.4),
                                   Inches(9.2), Inches(0.35) * (len(rows) + 1)).table
    for i, h in enumerate(cols):
        cell = table.cell(0, i)
        cell.text = h
        cell.text_frame.paragraphs[0].font.size = Pt(12)
        cell.text_frame.paragraphs[0].font.bold = True
    for ri, r in enumerate(rows, 1):
        vals = [f"{r.kpi_name} ({r.unit})" if r.unit else r.kpi_name, _fmt(r.actual), _fmt(r.target),
                _fmt(r.variance_pct, True), _fmt(r.change_vs_prior_pct, True), r.rag]
        for ci, v in enumerate(vals):
            cell = table.cell(ri, ci)
            cell.text = str(v)
            cell.text_frame.paragraphs[0].font.size = Pt(11)
        rag_cell = table.cell(ri, 5)
        rag_cell.fill.solid()
        rag_cell.fill.fore_color.rgb = RGBColor.from_string(RAG_FILL.get(r.rag, "FFFFFF"))
    table.columns[0].width = Inches(3.6)
    for i in range(1, 6):
        table.columns[i].width = Inches(1.12)


def build_pptx(period, rows) -> bytes:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = f"MIS Pack - {period.label}"
    s.placeholders[1].text = "KPI performance vs target and prior period"

    def chunks(items, n=10):
        for i in range(0, len(items), n):
            yield items[i:i + n]

    group = [r for r in rows if r.business_unit == "GROUP"]
    for i, part in enumerate(chunks(group)):
        _table_slide(prs, "Group KPI Summary" + (f" ({i + 1})" if i else ""), part)
    exc = [r for r in rows if r.rag in ("RED", "AMBER")]
    for i, part in enumerate(chunks(exc)):
        _table_slide(prs, "Exceptions by Unit" + (f" ({i + 1})" if i else ""),
                     [_Labeled(r) for r in part])
    for bu in sorted({r.business_unit for r in rows if r.business_unit != "GROUP"}):
        for i, part in enumerate(chunks([r for r in rows if r.business_unit == bu])):
            _table_slide(prs, f"{bu} KPIs" + (f" ({i + 1})" if i else ""), part)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


class _Labeled:
    """Prefix KPI name with unit code for exception slides."""
    def __init__(self, r):
        self._r = r

    def __getattr__(self, item):
        if item == "kpi_name":
            return f"[{self._r.business_unit}] {self._r.kpi_name}"
        return getattr(self._r, item)
